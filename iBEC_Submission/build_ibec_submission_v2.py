#!/usr/bin/env python
"""Build the competition-focused iBEC preliminary submission package."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(__file__).resolve().parent
BUILD = SOURCE / "build_v2"
PACKAGE = BUILD / "AnalyzingDisorder"
ACCENT = colors.HexColor("#087F6D")
DARK = colors.HexColor("#17212B")
MUTED = colors.HexColor("#52606D")
PALE = colors.HexColor("#E9F5F2")
WARM = colors.HexColor("#FFF4D6")


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def register_fonts():
    regular = Path("C:/Windows/Fonts/msyh.ttc")
    bold = Path("C:/Windows/Fonts/msyhbd.ttc")
    if regular.exists():
        pdfmetrics.registerFont(TTFont("IBEC", str(regular), subfontIndex=0))
    if bold.exists():
        pdfmetrics.registerFont(TTFont("IBEC-Bold", str(bold), subfontIndex=0))
    return ("IBEC" if regular.exists() else "Helvetica",
            "IBEC-Bold" if bold.exists() else "Helvetica-Bold")


def markdown_blocks(text):
    lines = text.splitlines()
    blocks = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            table = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in row):
                    table.append(row)
                index += 1
            blocks.append(("table", table))
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append((f"h{level}", line[level:].strip()))
            index += 1
            continue
        if re.match(r"^\d+\.\s", line) or line.startswith("- "):
            blocks.append(("bullet", re.sub(r"^(?:\d+\.|-)\s+", "", line)))
            index += 1
            continue
        if line.startswith("`") and line.endswith("`"):
            blocks.append(("code", line.strip("`")))
            index += 1
            continue
        paragraph = [line]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if (not next_line or next_line.startswith(("#", "|", "- ", "`"))
                    or re.match(r"^\d+\.\s", next_line)):
                break
            paragraph.append(next_line)
            index += 1
        blocks.append(("p", " ".join(paragraph)))
    return blocks


def inline_markup(value):
    value = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    value = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", value)
    value = re.sub(r"`(.+?)`", r"<font name='Courier'>\1</font>", value)
    return value


def build_report_pdf(markdown_path, output_path):
    regular, bold = register_fonts()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        "IBECTitle", fontName=bold, fontSize=20, leading=27, alignment=TA_CENTER,
        textColor=DARK, spaceAfter=12,
    ))
    styles.add(ParagraphStyle(
        "IBECH1", fontName=bold, fontSize=14, leading=19, textColor=DARK,
        spaceBefore=12, spaceAfter=7,
    ))
    styles.add(ParagraphStyle(
        "IBECH2", fontName=bold, fontSize=11.5, leading=16, textColor=ACCENT,
        spaceBefore=9, spaceAfter=5,
    ))
    styles.add(ParagraphStyle(
        "IBECBody", fontName=regular, fontSize=9.2, leading=13.5,
        alignment=TA_JUSTIFY, textColor=DARK, spaceAfter=5,
    ))
    styles.add(ParagraphStyle(
        "IBECBullet", fontName=regular, fontSize=9, leading=13, leftIndent=14,
        firstLineIndent=-8, textColor=DARK, spaceAfter=3,
    ))
    styles.add(ParagraphStyle(
        "IBECCode", fontName="Courier", fontSize=8.2, leading=11,
        leftIndent=12, backColor=colors.HexColor("#F2F4F5"), spaceAfter=5,
    ))
    doc = SimpleDocTemplate(
        str(output_path), pagesize=A4, leftMargin=2.0 * cm, rightMargin=2.0 * cm,
        topMargin=1.8 * cm, bottomMargin=1.8 * cm,
        title="DisorderFlow iBEC Project Report",
    )
    story = []
    first_title = True
    for kind, value in markdown_blocks(markdown_path.read_text(encoding="utf-8")):
        if kind == "h1" and first_title:
            story.extend([Spacer(1, 2.0 * cm), Paragraph(inline_markup(value), styles["IBECTitle"]),
                          Spacer(1, 0.5 * cm)])
            first_title = False
        elif kind in {"h1", "h2"}:
            story.append(Paragraph(inline_markup(value), styles["IBECH1"]))
        elif kind == "h3":
            story.append(Paragraph(inline_markup(value), styles["IBECH2"]))
        elif kind == "p":
            story.append(Paragraph(inline_markup(value), styles["IBECBody"]))
        elif kind == "bullet":
            story.append(Paragraph("• " + inline_markup(value), styles["IBECBullet"]))
        elif kind == "code":
            story.append(Paragraph(inline_markup(value), styles["IBECCode"]))
        elif kind == "table" and value:
            cells = [[Paragraph(inline_markup(cell), styles["IBECBody"]) for cell in row]
                     for row in value]
            table = Table(cells, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), DARK),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CAD2D8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8F7")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([table, Spacer(1, 5)])
    doc.build(story)


SLIDES = [
    {
        "title": "DisorderFlow",
        "subtitle": "An Auditable Platform for Antibody Design against Disordered Epitopes",
        "bullets": ["Team 28 - Analyzing Disorder", "Qilu University of Technology",
                    "AI-driven Life Science Discovery | iBEC 2026"],
    },
    {
        "title": "1. Why Disordered Epitopes?",
        "subtitle": "Rigid-pose antibody design is poorly matched to conformational heterogeneity",
        "bullets": [
            "Aβ, tau and α-synuclein are central protein-misfolding targets",
            "One antigen sequence can occupy multiple structural states",
            "A candidate should not be selected by one pose or one uncalibrated confidence score",
            "Engineering goal: auditable generation, counterfactual controls and uncertainty-aware filtering",
        ],
    },
    {
        "title": "2. Original Contributions",
        "subtitle": "A platform, not a single opaque score",
        "bullets": [
            "ECLS: complex versus peptide-stripped structural likelihood contrast",
            "Position-sensitive antigen-disorder routing in a Bayesian Flow Network",
            "BFN expands H3 exploration, with an explicitly measured developability tradeoff",
            "Fail-closed binder confidence: abstain when experimental labels are absent",
        ],
    },
    {
        "title": "3. End-to-End Platform",
        "subtitle": "Structure → generation → multi-state scoring → filtering → validation → handoff",
        "bullets": [
            "BFN, ProteinMPNN and ESM-IF matched workflows",
            "Gradio interface plus reproducible command-line runners",
            "AF2-Multimer through a persistent WSL2 GPU worker",
            "Developability, diversity, provenance and claim-boundary exports",
        ],
    },
    {
        "title": "4. Frozen ECLS Temporal Result",
        "subtitle": "31 structures | 15 antigen clusters | 200 shuffles per record",
        "bullets": [
            "Mean native advantage: 0.1723",
            "Cluster-bootstrap 95% CI: [0.0592, 0.2919]",
            "Positive clusters: 12/15 (80%)",
            "Bounded claim: structural sequence contrast, not affinity",
        ],
    },
    {
        "title": "5. Aβ Candidate Engineering Funnel",
        "subtitle": "A frozen, auditable 3D6/4HIX case study",
        "bullets": [
            "481 source memberships → 352 unique hypotheses",
            "77 pass computational gates → 24 diverse candidates",
            "27 entities × 3 AF2 seeds = 81/81 successful predictions",
            "16 pass AF2/PRODIGY gates → 12 final computational candidates",
        ],
    },
    {
        "title": "6. Multi-Seed Structural Validation",
        "subtitle": "Small score changes are reported with uncertainty",
        "bullets": [
            "Native median ipTM: 0.4762",
            "Candidate median ipTM range: 0.4763-0.5034",
            "Candidate seed ranges: 0.0030-0.0087",
            "Interface PAE ≈ 26: candidates remain computational hypotheses",
        ],
    },
    {
        "title": "7. Evidence Audit and Second Target",
        "subtitle": "Correcting claims is part of the engineering workflow",
        "bullets": [
            "1,289 records × 5 conformations; 219 have exact DisProt disorder-region evidence",
            "The unsupported phrase '800+ natural IDPs' has been removed",
            "Tau/5MP3: correct 13-aa H3, five poses, 275 unique hypotheses, 163 pass all gates",
            "Eight diverse Tau candidates; minimum pairwise Hamming distance 2",
        ],
    },
    {
        "title": "8. Engineering and Translation",
        "subtitle": "Computational output is converted into an auditable handoff",
        "bullets": [
            "Secreted scFv contract: signal-VH-(G4S)3-VL-G4S-His6",
            "12 candidates plus native and counterfactual sequence controls",
            "Blinded assay layouts and measurement import schemas",
            "No false probability: binder-confidence status is abstain_not_trained",
        ],
    },
    {
        "title": "9. iBEC Deliverables",
        "subtitle": "Primary track: AI-driven Life Science Discovery",
        "bullets": [
            "Runnable platform and 4HIX demonstration workflow",
            "Frozen ECLS benchmark and 449-artifact release lineage",
            "Aβ 352→24→12 case study with 81 hashed AF2 structures",
            "Completed audits and Tau 275→163→8 second-target extension",
            "Current boundary: no experimental binding, affinity or therapeutic claim",
        ],
    },
]


def add_ppt_text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    paragraph.alignment = align
    return box


def build_abstract_pptx(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    dark = RGBColor(23, 33, 43)
    accent = RGBColor(8, 127, 109)
    muted = RGBColor(82, 96, 109)
    pale = RGBColor(233, 245, 242)
    for index, content in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(248, 250, 249)
        bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.16), Inches(7.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()
        if index == 1:
            add_ppt_text(slide, 0.8, 1.2, 11.8, 0.8, content["title"], 38, dark, True,
                         PP_ALIGN.CENTER)
            add_ppt_text(slide, 1.0, 2.15, 11.4, 0.7, content["subtitle"], 20, muted,
                         False, PP_ALIGN.CENTER)
            y = 3.5
            for bullet in content["bullets"]:
                add_ppt_text(slide, 2.0, y, 9.3, 0.42, bullet, 14, accent, index == 1,
                             PP_ALIGN.CENTER)
                y += 0.55
        else:
            add_ppt_text(slide, 0.65, 0.45, 11.8, 0.55, content["title"], 26, dark, True)
            add_ppt_text(slide, 0.68, 1.1, 11.7, 0.45, content["subtitle"], 14, accent, True)
            y = 1.9
            for bullet in content["bullets"]:
                shape = slide.shapes.add_shape(5, Inches(0.8), Inches(y), Inches(11.55), Inches(0.72))
                shape.fill.solid()
                shape.fill.fore_color.rgb = pale
                shape.line.color.rgb = RGBColor(191, 220, 213)
                add_ppt_text(slide, 1.05, y + 0.12, 11.0, 0.45, "•  " + bullet, 15, dark)
                y += 0.9
        add_ppt_text(slide, 0.7, 7.05, 5.5, 0.25, "Team 28 | DisorderFlow | iBEC 2026", 8,
                     muted)
        add_ppt_text(slide, 11.8, 7.05, 0.8, 0.25, f"{index}/10", 8, muted, False,
                     PP_ALIGN.RIGHT)
    prs.save(output_path)


def build_abstract_pdf(output_path):
    regular, bold = register_fonts()
    page_w, page_h = 13.333 * 72, 7.5 * 72
    pdf = canvas.Canvas(str(output_path), pagesize=(page_w, page_h))
    for index, content in enumerate(SLIDES, 1):
        pdf.setFillColor(colors.HexColor("#F8FAF9"))
        pdf.rect(0, 0, page_w, page_h, 0, 1)
        pdf.setFillColor(ACCENT)
        pdf.rect(0, 0, 9, page_h, 0, 1)
        if index == 1:
            pdf.setFont(bold, 34)
            pdf.setFillColor(DARK)
            pdf.drawCentredString(page_w / 2, page_h - 145, content["title"])
            pdf.setFont(regular, 16)
            pdf.setFillColor(MUTED)
            pdf.drawCentredString(page_w / 2, page_h - 180, content["subtitle"])
            y = page_h - 250
            for bullet in content["bullets"]:
                pdf.setFont(regular, 12)
                pdf.setFillColor(ACCENT)
                pdf.drawCentredString(page_w / 2, y, bullet)
                y -= 28
        else:
            pdf.setFont(bold, 23)
            pdf.setFillColor(DARK)
            pdf.drawString(40, page_h - 55, content["title"])
            pdf.setFont(bold, 12)
            pdf.setFillColor(ACCENT)
            pdf.drawString(42, page_h - 82, content["subtitle"])
            y = page_h - 135
            for bullet in content["bullets"]:
                pdf.setFillColor(PALE)
                pdf.roundRect(48, y - 10, page_w - 96, 45, 5, 0, 1)
                pdf.setFont(regular, 12)
                pdf.setFillColor(DARK)
                pdf.drawString(65, y + 8, "•  " + bullet)
                y -= 58
        pdf.setFont(regular, 8)
        pdf.setFillColor(MUTED)
        pdf.drawString(40, 20, "Team 28 | DisorderFlow | iBEC 2026")
        pdf.drawRightString(page_w - 40, 20, f"{index}/10")
        pdf.showPage()
    pdf.save()


def write_lightweight_artifacts(supplement):
    supplement.mkdir(parents=True, exist_ok=True)
    sources = {
        "ecls_final_decision.json": ROOT / "results/publication/h3_ecls_temporal_final_v1/final_decision.json",
        "abeta_candidate_summary.json": ROOT / "results/prospective/abeta_4hix_af2_gate_v1/analysis.json",
        "abeta_final_shortlist.csv": ROOT / "results/prospective/abeta_4hix_af2_gate_v1/shortlist_for_expression_prep.csv",
        "binder_confidence_status.json": ROOT / "results/prospective/binder_confidence_v1/measurement_status.json",
        "e1_generator_behavior_summary.json": ROOT / "results/ibec/e1_generator_behavior_v1/summary.json",
        "e2_idp_evidence_summary.json": ROOT / "results/ibec/e2_idp_evidence_audit_v1/summary.json",
        "tau_sanity_status.json": ROOT / "results/ibec/tau_sanity_v1/status.json",
        "tau_candidate_results.json": ROOT / "results/ibec/tau_candidates_v1/results.json",
        "REPRODUCIBILITY.md": ROOT / "publication/REPRODUCIBILITY.md",
    }
    for name, source in sources.items():
        if name == "abeta_candidate_summary.json":
            data = json.loads(source.read_text(encoding="utf-8"))
            compact = {
                "schema_version": data["schema_version"],
                "status": data["status"],
                "baselines": data["baselines"],
                "summary": data["summary"],
                "shortlist": [
                    {key: row[key] for key in (
                        "final_rank", "entity_id", "h3_sequence", "median_iptm",
                        "worst_iptm", "iptm_seed_range", "median_interface_pae",
                        "median_prodigy_delta_g_kcal_mol",
                    )}
                    for row in data["shortlist"]
                ],
                "claim_boundary": data["claim_boundary"],
            }
            (supplement / name).write_text(json.dumps(compact, indent=2) + "\n", encoding="ascii")
        elif name == "binder_confidence_status.json":
            data = json.loads(source.read_text(encoding="utf-8"))
            compact = {
                "schema_version": data["schema_version"],
                "status": data["status"],
                "assay_validity": data["assay_validity"],
                "complete_candidate_labels": data["complete_candidate_labels"],
                "training_readiness": data["training_readiness"],
                "probability_output": data["probability_output"],
                "probability_status": data["probability_status"],
                "claim_boundary": data["claim_boundary"],
                "privacy_note": "Blind sample mappings and the blinding key are excluded.",
            }
            (supplement / name).write_text(json.dumps(compact, indent=2) + "\n", encoding="ascii")
        elif name == "tau_candidate_results.json":
            data = json.loads(source.read_text(encoding="ascii"))
            compact = {
                "schema_version": data["schema_version"],
                "status": data["status"],
                "target": data["target"],
                "reference": data["reference"],
                "summary": data["summary"],
                "candidate_source_memberships": data["candidate_source_memberships"],
                "shortlist": [
                    {key: row[key] for key in (
                        "rank", "sequence", "sources", "mutation_count",
                        "bfn_mean_cms", "bfn_worst_cms", "bfn_std_cms",
                        "contact_mean_delta_native", "contact_min_delta_native",
                        "full_heavy_developability_risk", "risk_increase_over_native",
                        "gate_checks", "all_gates_pass",
                    )}
                    for row in data["shortlist"]
                ],
                "claim_boundary": data["claim_boundary"],
            }
            (supplement / name).write_text(json.dumps(compact, indent=2) + "\n", encoding="ascii")
        else:
            shutil.copy2(source, supplement / name)
    shutil.copy2(SOURCE / "SUPPLEMENTARY_README_V2.md", supplement / "README.md")
    shutil.copy2(SOURCE / "IBEC_SCORING_MAP_V2.md", supplement / "IBEC_SCORING_MAP.md")
    shutil.copy2(SOURCE / "IBEC_TRACK_DECISION.md", supplement / "IBEC_TRACK_DECISION.md")
    manifest = []
    for path in sorted(supplement.iterdir()):
        if path.is_file():
            manifest.append({"path": path.name, "bytes": path.stat().st_size, "sha256": sha256(path)})
    (supplement / "MANIFEST.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="ascii")


def main():
    if BUILD.exists():
        shutil.rmtree(BUILD)
    PACKAGE.mkdir(parents=True)
    report_pdf = PACKAGE / "28-AnalyzingDisorder-ProjectReport.pdf"
    abstract_pdf = PACKAGE / "28-AnalyzingDisorder-TechnicalAbstract.pdf"
    abstract_pptx = PACKAGE / "28-AnalyzingDisorder-TechnicalAbstract.pptx"
    build_report_pdf(SOURCE / "iBEC_PROJECT_REPORT_V2.md", report_pdf)
    build_abstract_pdf(abstract_pdf)
    build_abstract_pptx(abstract_pptx)
    write_lightweight_artifacts(PACKAGE / "Supplementary_Materials")
    shutil.copy2(SOURCE / "SUBMISSION_CHECKLIST_V2.md", PACKAGE / "SUBMISSION_CHECKLIST.md")
    logo = SOURCE / "logo/28-解析无序-logo.png"
    if logo.exists():
        shutil.copy2(logo, PACKAGE / "28-AnalyzingDisorder-logo.png")
    package_manifest = []
    for path in sorted(PACKAGE.rglob("*")):
        if path.is_file():
            package_manifest.append({
                "path": path.relative_to(PACKAGE).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            })
    (PACKAGE / "PACKAGE_MANIFEST.json").write_text(
        json.dumps(package_manifest, indent=2) + "\n", encoding="ascii")
    archive = BUILD / "AnalyzingDisorder.zip"
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as handle:
        for path in sorted(PACKAGE.rglob("*")):
            if path.is_file():
                handle.write(path, Path("AnalyzingDisorder") / path.relative_to(PACKAGE))
    summary = {
        "status": "ibec_preliminary_package_built",
        "archive": str(archive.relative_to(ROOT).as_posix()),
        "archive_bytes": archive.stat().st_size,
        "archive_megabytes": round(archive.stat().st_size / 1024 / 1024, 3),
        "under_50_mb": archive.stat().st_size <= 50 * 1024 * 1024,
        "archive_sha256": sha256(archive),
        "required_files": {
            "project_report_pdf": report_pdf.exists(),
            "technical_abstract_pdf": abstract_pdf.exists(),
        },
        "technical_abstract_pages": len(SLIDES),
        "submission_deadline": "2026-09-30",
    }
    (BUILD / "BUILD_SUMMARY.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="ascii")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
