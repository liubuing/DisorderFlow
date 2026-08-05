"""
iBEC Technical Abstract PPT Generator
10 slides, 16:9 format, to be exported as PDF.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Colors
BG_WHITE = RGBColor(0xF7, 0xF8, 0xF6)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x08, 0x7F, 0x6D)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GRAY = RGBColor(0x99, 0x99, 0x99)
TABLE_HEADER = RGBColor(0x2D, 0x34, 0x36)
TABLE_ALT = RGBColor(0xF0, 0xF3, 0xF5)
GREEN_HL = RGBColor(0xD4, 0xED, 0xDA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text_box(slide, left, top, width, height, text, font_size=14,
                 bold=False, color=DARK_TEXT, alignment=PP_ALIGN.LEFT,
                 font_name='Calibri'):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                     Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_bullet_slide(slide, left, top, width, height, items, font_size=13,
                     color=DARK_TEXT, bullet_char='\u2022'):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                     Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"{bullet_char} {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = 'Calibri'
        p.space_after = Pt(4)
    return txBox

def add_table(slide, left, top, rows, cols, data, col_widths=None):
    """Add a styled table to the slide."""
    tbl_shape = slide.shapes.add_table(rows, cols,
                                        Inches(left), Inches(top),
                                        Inches(col_widths[0] if col_widths else sum([2.0]*cols)),
                                        Inches(0.35 * rows))
    tbl = tbl_shape.table

    # Set column widths
    if col_widths:
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = Inches(w)

    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = str(data[r][c])
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(11)
                paragraph.font.name = 'Calibri'
                if r == 0:
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = WHITE
                else:
                    paragraph.font.color.rgb = DARK_TEXT

            # Cell fill
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_HEADER
            elif r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_ALT
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE

    return tbl_shape

def add_section_header(slide, number, title):
    """Add a consistent section header with accent bar."""
    # Accent bar
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                    Inches(0.5), Inches(0.4),
                                    Inches(0.06), Inches(0.45))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()

    # Section number + title
    add_text_box(slide, 0.7, 0.35, 8, 0.5,
                 f"{number}. {title}", font_size=22, bold=True, color=DARK_TEXT)

def add_footer(slide, page_num):
    """Add footer with page number."""
    add_text_box(slide, 8.5, 5.2, 1.2, 0.3,
                 f"{page_num} / 10", font_size=9, color=LIGHT_GRAY,
                 alignment=PP_ALIGN.RIGHT)
    add_text_box(slide, 0.5, 5.2, 3, 0.3,
                 "Team 28 — Analyzing Disorder | iBEC 2026",
                 font_size=9, color=LIGHT_GRAY)


def build_pptx():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    blank_layout = prs.slide_layouts[6]  # blank

    # ==================== SLIDE 1: TITLE ====================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_WHITE)

    # Accent line
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                    Inches(2), Inches(1.5),
                                    Inches(6), Inches(0.04))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()

    add_text_box(slide, 1, 1.7, 8, 1.0,
                 "DisorderFlow", font_size=36, bold=True,
                 color=DARK_TEXT, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, 1, 2.5, 8, 0.6,
                 "Bayesian Flow Networks for Disorder-Aware Antibody CDR-H3 Design",
                 font_size=16, color=GRAY, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, 1, 3.4, 8, 0.4,
                 "Team 28 — Analyzing Disorder",
                 font_size=14, color=ACCENT, alignment=PP_ALIGN.CENTER, bold=True)
    add_text_box(slide, 1, 3.8, 8, 0.3,
                 "Chen Haoyang, Hu Jing, Yu Jiarui  |  Qilu University of Technology",
                 font_size=12, color=GRAY, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, 1, 4.2, 8, 0.3,
                 "iBEC 2026 — International Bioinformatics Engineering Competition",
                 font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # ==================== SLIDE 2: PROBLEM ====================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_WHITE)
    add_section_header(slide, 1, "The IDP Antibody Design Challenge")
    add_footer(slide, 2)

    add_bullet_slide(slide, 0.7, 1.1, 8.5, 4.0, [
        "Intrinsically disordered proteins (IDPs) lack stable 3D structure yet drive neurodegenerative diseases",
        "Therapeutic targets: Amyloid-beta (Alzheimer's), tau, alpha-synuclein (Parkinson's)",
        "Most antibody design methods assume rigid, well-ordered epitopes — inadequate for IDPs",
        "CDR-H3 is the primary determinant of antigen specificity and the most diverse loop",
        "Key gap: no computational framework couples epitope disorder profiles to CDR-H3 sequence selection",
    ], font_size=14)

    # Bottom callout
    add_text_box(slide, 0.7, 4.3, 8.5, 0.5,
                 "Goal: Build a disorder-aware scoring framework that conditions CDR-H3 design on "
                 "both backbone geometry and per-residue epitope disorder profiles.",
                 font_size=12, bold=True, color=ACCENT)

    # ==================== SLIDE 3: PLATFORM ====================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_WHITE)
    add_section_header(slide, 2, "DisorderFlow Platform Architecture")
    add_footer(slide, 3)

    # Four components as boxes
    components = [
        ("BFN Core", "Categorical distributions\nover 20 amino acids\nIterative refinement"),
        ("Disorder Head", "Per-residue disorder\nprediction (focal loss\ngamma=4.0)"),
        ("Position Routing", "Couple disorder values\nto CDR-antigen pair\nfeatures directly"),
        ("Contrastive Loss", "Factual vs shuffled/\nmismatched profiles\nRank-aware training"),
    ]
    for i, (title, desc) in enumerate(components):
        x = 0.5 + i * 2.35
        # Box
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                        Inches(x), Inches(1.2),
                                        Inches(2.1), Inches(1.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xE8, 0xF5, 0xF3)
        shape.line.color.rgb = ACCENT
        shape.line.width = Pt(1.5)

        add_text_box(slide, x + 0.1, 1.3, 1.9, 0.4,
                     title, font_size=13, bold=True, color=ACCENT,
                     alignment=PP_ALIGN.CENTER)
        add_text_box(slide, x + 0.1, 1.7, 1.9, 1.2,
                     desc, font_size=10, color=GRAY,
                     alignment=PP_ALIGN.CENTER)

    # Training pipeline description
    add_text_box(slide, 0.7, 3.3, 8.5, 0.4,
                 "Training Pipeline", font_size=14, bold=True, color=DARK_TEXT)
    add_text_box(slide, 0.7, 3.7, 8.5, 1.0,
                 "SAbDab2 structural database  |  6-axis homology-disjoint split (PDB, paired Ab, VH, VL, H3, antigen)  |  "
                 "Backbone coordinates + disorder labels + negative profiles  |  ProteinMPNN v_48_020 for evaluation",
                 font_size=11, color=GRAY)

    # ==================== SLIDE 4: ECLS METRIC ====================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_WHITE)
    add_section_header(slide, 3, "ECLS: Epitope-Conditioned Likelihood Shift")
    add_footer(slide, 4)

    # Formula box
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(1.5), Inches(1.2),
                                    Inches(7), Inches(0.8))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xF0, 0xF3, 0xF5)
    shape.line.color.rgb = RGBColor(0xDE, 0xE2, 0xE6)
    add_text_box(slide, 1.7, 1.3, 6.6, 0.6,
                 "ECLS(s) = NLL_complex(s) - NLL_peptide-stripped(s)",
                 font_size=18, bold=True, color=DARK_TEXT, alignment=PP_ALIGN.CENTER,
                 font_name='Consolas')

    add_text_box(slide, 0.7, 2.2, 8.5, 0.8,
                 "Measures whether a CDR-H3 sequence is preferentially supported by the presence of "
                 "peptide epitope coordinates. Positive advantage = native sequence benefits more from "
                 "epitope context than composition-matched shuffles.",
                 font_size=12, color=GRAY)

    add_text_box(slide, 0.7, 3.0, 8.5, 0.4,
                 "Methodology", font_size=14, bold=True, color=DARK_TEXT)
    add_bullet_slide(slide, 0.7, 3.4, 8.5, 2.0, [
        "Score only heavy chain; fix non-H3 positions; retain light chain + peptide context",
        "200 composition-matched shuffles per record as counterfactual controls",
        "Official SAbDab2 CDR-H3 annotations; heavy-atom distance <= 4.5 A contact requirement",
        "Official antigen cluster as inference unit; 10,000 bootstrap resamples for CIs",
    ], font_size=12)

    # ==================== SLIDE 5: ECLS RESULTS ====================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_WHITE)
    add_section_header(slide, 4, "Results: ECLS Detection and Temporal Transfer")
    add_footer(slide, 5)

    # Adaptation table
    add_text_box(slide, 0.7, 1.1, 4, 0.3,
                 "Adaptation Set (46 clusters)", font_size=12, bold=True, color=DARK_TEXT)
    data1 = [
        ['Metric', 'Value'],
        ['Mean ECLS advantage', '0.217'],
        ['Median', '0.203'],
        ['95% Bootstrap CI', '[0.146, 0.290]'],
        ['Positive clusters', '80.4%'],
    ]
    add_table(slide, 0.7, 1.45, 5, 2, data1, col_widths=[2.2, 1.5])

    # Temporal final table
    add_text_box(slide, 5.0, 1.1, 4.5, 0.3,
                 "Temporal Final (15 clusters, post-2021)", font_size=12, bold=True, color=DARK_TEXT)
    data2 = [
        ['Metric', 'Value'],
        ['Mean ECLS advantage', '0.172'],
        ['Median', '0.119'],
        ['95% Bootstrap CI', '[0.059, 0.292]'],
        ['Positive clusters', '80.0%'],
    ]
    add_table(slide, 5.0, 1.45, 5, 2, data2, col_widths=[2.2, 1.5])

    # Summary
    add_text_box(slide, 0.7, 3.6, 8.5, 0.8,
                 "Both frozen gate sets passed. Deposited CDR-H3 sequences carry measurable "
                 "epitope-conformation-specific signal that transfers to post-training structures. "
                 "Sign-flip P < 1e-6 (adaptation) and P = 0.014 (temporal final).",
                 font_size=12, color=GRAY)

    # Key finding box
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(0.7), Inches(4.4),
                                    Inches(8.5), Inches(0.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GREEN_HL
    shape.line.fill.background()
    add_text_box(slide, 0.9, 4.45, 8.1, 0.4,
                 "Key: 80%+ positive clusters on sealed temporal data confirms non-trivial epitope-H3 coupling",
                 font_size=11, bold=True, color=RGBColor(0x15, 0x57, 0x24))

    # ==================== SLIDE 6: T2.1 v2 ====================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_WHITE)
    add_section_header(slide, 5, "T2.1 v2: Deterministic Torsion Recovery Benchmark")
    add_footer(slide, 6)

    add_text_box(slide, 0.7, 1.1, 8.5, 0.6,
                 "Deterministic phi/psi torsion perturbation to 1.5-3.0 A RMSD tier, then test if native "
                 "residue contacts guide recovery. Peptide restraint: 100 kJ/mol/nm2.",
                 font_size=12, color=GRAY)

    data3 = [
        ['Metric', 'Value'],
        ['Total structures', '31'],
        ['Torsion hits (target tier)', '26/31 (83.9%)'],
        ['Valid structures after recovery', '25/26 (96.2%)'],
        ['Mean held-out contact recovery', '0.581'],
        ['95% CI', '[0.533, 0.635]'],
        ['Positive recovery fraction', '100%'],
    ]
    add_table(slide, 0.7, 1.8, 7, 2, data3, col_widths=[2.5, 2.0])

    # Comparison box
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(5.5), Inches(1.8),
                                    Inches(4.0), Inches(2.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xF8, 0xF9, 0xFA)
    shape.line.color.rgb = RGBColor(0xDE, 0xE2, 0xE6)

    add_text_box(slide, 5.7, 1.9, 3.6, 0.3,
                 "T2 (failed) vs T2.1 v2 (success)", font_size=11, bold=True, color=DARK_TEXT)
    add_text_box(slide, 5.7, 2.3, 3.6, 1.5,
                 "T2:  3/7 valid, recovery = -0.017\n"
                 "       (high-temp perturbation, poor calibration)\n\n"
                 "T2.1 v2:  25/26 valid, recovery = 0.581\n"
                 "       (deterministic torsion, 100 kJ/mol/nm2)\n\n"
                 "Root cause: parameter calibration, not\n"
                 "fundamental method limitation.",
                 font_size=10, color=GRAY)

    # Bottom highlight
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(0.7), Inches(4.5),
                                    Inches(8.5), Inches(0.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GREEN_HL
    shape.line.fill.background()
    add_text_box(slide, 0.9, 4.55, 8.1, 0.4,
                 "Gate pass: 96.2% valid structures, 100% positive contact recovery, mean recovery 0.581",
                 font_size=11, bold=True, color=RGBColor(0x15, 0x57, 0x24))

    # ==================== SLIDE 7: 4HIX DESIGN ====================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_WHITE)
    add_section_header(slide, 6, "4HIX IDP Antibody Design Validation")
    add_footer(slide, 7)

    add_text_box(slide, 0.7, 1.1, 8.5, 0.5,
                 "Scaffold: 4HIX (humanized 3D6 Fab, Abeta 1-6 epitope DAEFRH). "
                 "ProteinMPNN T=0.5, 20 samples. AF2 multimer V3 with VH:VL chain separator.",
                 font_size=12, color=GRAY)

    data4 = [
        ['Rank', 'H3 Sequence', 'ipTM', 'pLDDT', 'iPAE'],
        ['Native', 'VRYDHYSGSSDY', '0.449', '0.194', '24.7'],
        ['#1', 'LYDESKDAESE', '0.462', '0.200', '24.5'],
        ['#2', 'LYDAHHGAHSL', '0.461', '0.192', '24.4'],
        ['#3', 'LYDGSIGAESQ', '0.460', '0.195', '24.5'],
        ['#4', 'LYDSSVDASGH', '0.459', '0.199', '24.7'],
        ['#5', 'LFNEANCAESY', '0.458', '0.199', '24.3'],
    ]
    tbl_shape = add_table(slide, 0.7, 1.7, 7, 5, data4,
                          col_widths=[0.6, 1.8, 0.7, 0.7, 0.7])

    # Highlight top design row (row index 2 = #1)
    tbl = tbl_shape.table
    for c in range(5):
        cell = tbl.cell(2, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = GREEN_HL

    # Key findings
    add_text_box(slide, 5.2, 1.7, 4.3, 0.3,
                 "Key Findings", font_size=13, bold=True, color=DARK_TEXT)
    add_bullet_slide(slide, 5.2, 2.1, 4.3, 2.5, [
        "20/20 designs pass AF2 validation",
        "Top 3 exceed native ipTM (0.449)",
        "Best design: ipTM = 0.462 (+2.9%)",
        "AF2 chain separator fix: VH:VL vs concatenated prevents 4.5x ipTM underestimation",
    ], font_size=11)

    # Bottom note
    add_text_box(slide, 0.7, 4.6, 8.5, 0.4,
                 "Note: 6-residue epitope; AF2 metrics less discriminative for short peptides. "
                 "4HIX was in AF2 training set (2012). No wet-lab validation performed.",
                 font_size=10, color=LIGHT_GRAY)

    # ==================== SLIDE 8: EVIDENCE HIERARCHY ====================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_WHITE)
    add_section_header(slide, 7, "Evidence Hierarchy and Methodological Rigor")
    add_footer(slide, 8)

    evidence = [
        ("1. ECLS Temporal Final", "n=15 clusters, post-2021, sealed", "Positive discrimination (0.172, 80%+)"),
        ("2. ECLS Adaptation", "n=46 clusters, exposed", "Replicates temporal final pattern"),
        ("3. T2.1 v2 Recovery", "n=31 structures, 96.2% valid", "Physical contact recovery (0.581)"),
        ("4. 4HIX Design", "n=20 designs, 20/20 AF2 pass", "Prospective design (top ipTM 0.462)"),
        ("5. T1 Ensemble", "n=7 clusters", "Single-pose overconfidence (-0.200)"),
        ("6. Generator Calibration", "n=7 clusters, exploratory", "Calibrated reranker works"),
    ]

    for i, (title, detail, result) in enumerate(evidence):
        y = 1.2 + i * 0.6
        # Strength indicator
        intensity = max(0.3, 1.0 - i * 0.12)
        bar_color = RGBColor(int(8 * intensity), int(127 * intensity), int(109 * intensity))
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                        Inches(0.7), Inches(y),
                                        Inches(0.08), Inches(0.4))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bar_color
        shape.line.fill.background()

        add_text_box(slide, 0.95, y - 0.02, 3.5, 0.3,
                     title, font_size=12, bold=True, color=DARK_TEXT)
        add_text_box(slide, 0.95, y + 0.25, 3.5, 0.25,
                     detail, font_size=9, color=LIGHT_GRAY)
        add_text_box(slide, 5.0, y + 0.05, 4.5, 0.3,
                     result, font_size=11, color=GRAY)

    # Rigor note
    add_text_box(slide, 0.7, 4.9, 8.5, 0.4,
                 "All evaluations use frozen gates, one-shot evaluation on sealed data, and "
                 "homology-disjoint splitting across 6 axes. Bootstrap CIs with 10,000 resamples.",
                 font_size=10, color=LIGHT_GRAY)

    # ==================== SLIDE 9: SOFTWARE & REPRODUCIBILITY ====================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_WHITE)
    add_section_header(slide, 8, "Software Implementation and Reproducibility")
    add_footer(slide, 9)

    add_text_box(slide, 0.7, 1.1, 4.0, 0.3,
                 "Software Components", font_size=13, bold=True, color=DARK_TEXT)
    add_bullet_slide(slide, 0.7, 1.5, 4.0, 2.5, [
        "BFN core: categorical distributions, focal-weighted losses",
        "Disorder-augmented dataset: label injection + contrastive profiles",
        "Receiver network: position-sensitive routing",
        "IDP design pipeline: 5-stage target-to-validation",
        "AF2 infrastructure: proper chain separator handling",
    ], font_size=11)

    add_text_box(slide, 5.2, 1.1, 4.3, 0.3,
                 "Reproducibility", font_size=13, bold=True, color=DARK_TEXT)
    add_bullet_slide(slide, 5.2, 1.5, 4.3, 2.5, [
        "Frozen YAML configuration contracts",
        "SHA256 manifests for all result files",
        "results_t2.1_v2_final.json (31 records)",
        "4hix_final_validation/final_report.json (20 designs)",
        "INTEGRATIVE_ANALYSIS.md evidence chain",
    ], font_size=11)

    # Key bugs fixed box
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(0.7), Inches(3.7),
                                    Inches(8.5), Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xFF, 0xF3, 0xCD)
    shape.line.color.rgb = RGBColor(0xFF, 0xC1, 0x07)

    add_text_box(slide, 0.9, 3.8, 8.1, 0.3,
                 "Critical Bugs Fixed During Development", font_size=12, bold=True,
                 color=RGBColor(0x85, 0x6A, 0x04))
    add_text_box(slide, 0.9, 4.15, 8.1, 0.7,
                 "1. Disorder head: training batches never contained disorder_label -> PaddingCollate dropped the field -> loss never fired. "
                 "Fixed with explicit label injection.\n"
                 "2. AF2 chain separator: antibody chains must use VH:VL (not concatenation) or ipTM drops 4.5x.\n"
                 "3. T2.1 perturbation: peptide restraint 25 -> 100 kJ/mol/nm2 improved valid fraction from 43% to 96.2%.",
                 font_size=10, color=RGBColor(0x85, 0x6A, 0x04))

    # ==================== SLIDE 10: CONCLUSIONS ====================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_WHITE)
    add_section_header(slide, 9, "Conclusions and Future Work")
    add_footer(slide, 10)

    add_text_box(slide, 0.7, 1.1, 8.5, 0.3,
                 "Conclusions", font_size=14, bold=True, color=DARK_TEXT)
    add_bullet_slide(slide, 0.7, 1.5, 8.5, 1.5, [
        "ECLS establishes epitope-conditioned CDR-H3 signal across 46 clusters (mean 0.217, 80.4% positive)",
        "Temporal transfer confirmed on sealed post-2021 data (mean 0.172, 80% positive, P=0.014)",
        "T2.1 v2 achieves 96.2% valid structures with 100% positive contact recovery (0.581)",
        "4HIX prospective design: 20/20 pass AF2, top 3 exceed native ipTM",
        "Disorder-aware framework provides a principled approach to IDP antibody design",
    ], font_size=12)

    add_text_box(slide, 0.7, 3.2, 8.5, 0.3,
                 "Limitations", font_size=14, bold=True, color=DARK_TEXT)
    add_bullet_slide(slide, 0.7, 3.55, 8.5, 0.8, [
        "No wet-lab validation; all results computational",
        "Short epitope (6 residues); AF2 less discriminative; 4HIX in AF2 training set",
    ], font_size=11)

    add_text_box(slide, 0.7, 4.2, 8.5, 0.3,
                 "Future Work", font_size=14, bold=True, color=DARK_TEXT)
    add_bullet_slide(slide, 0.7, 4.55, 8.5, 0.8, [
        "CAID benchmark evaluation of retrained disorder head",
        "Extended IDP targets (tau, alpha-synuclein) and experimental validation (SPR/ELISA)",
    ], font_size=11)

    # Save
    output_path = r"C:\biological\DisorderFlow\iBEC_Submission\28-解析无序-技术摘要.pptx"
    prs.save(output_path)
    print(f"PPTX saved: {output_path}")
    return output_path

if __name__ == '__main__':
    build_pptx()
